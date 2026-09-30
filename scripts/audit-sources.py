"""Read-only source inventory. No workbook is ever saved."""
from pathlib import Path
from collections import Counter
import json, hashlib, zipfile, sys, xml.etree.ElementTree as ET
import openpyxl
sys.stdout.reconfigure(encoding='utf-8')

ROOT = Path(__file__).resolve().parents[1]
out = ROOT / 'reports'
out.mkdir(exist_ok=True)
report = {}
for path in sorted(ROOT.glob('*.xlsx')):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for ws in wb:
        rows = [(i, list(row)) for i, row in enumerate(ws.iter_rows(values_only=True), 1) if any(v is not None for v in row)]
        item = {'sheet': ws.title, 'rows_nonempty': len(rows), 'max_row': ws.max_row, 'max_column': ws.max_column, 'sample': rows[:5]}
        if path.name in ('Arbol_estructura_lamitex.xlsx', 'Datos_iniciales_lamitex.xlsx'):
            col = rows[0][1].index('colchon_id')
            item['rows_101'] = [r for r in rows[1:] if r[1][col] == 101]
        sheets.append(item)
    report[path.name] = {'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'sheets': sheets}
    wb.close()
for ext in (() if '--excel-only' in sys.argv else ('docx', 'pptx')):
    for path in ROOT.glob('*.' + ext):
        with zipfile.ZipFile(path) as z:
            entries = ['word/document.xml'] if ext == 'docx' else sorted(n for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml'))
            text = []
            for name in entries:
                tree = ET.fromstring(z.read(name))
                words = [e.text for e in tree.iter() if e.tag.endswith('}t') and e.text]
                text.append({'part': name, 'text': ' '.join(words)})
            report[path.name] = {'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'parts': text}
(out / 'source-audit.json').write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
for name, info in report.items():
    print(name)
    for s in info.get('sheets', []):
        print(json.dumps(s, ensure_ascii=False, default=str))
    for p in info.get('parts', [])[:2]:
        print(p['text'][:1600])
